
import os
import re
import sqlite3
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ==========================================
# RESTAURANT DATA ANALYST AGENT V2
# ==========================================

BASE = (
    "/storage/emulated/0/Download/"
    "Restaurant-Sales-Analytics"
)

DB_PATH = os.path.join(BASE, "restaurant_analytics.db")
REPORT_PATH = os.path.join(BASE, "AI_Data_Analyst_Report.xlsx")

TABLE = "sales"

REQUIRED_COLUMNS = {
    "Order_ID", "Order_Date", "Branch", "Product",
    "Category", "Quantity", "Unit_Price", "Discount",
    "Payment_Method", "Gross_Sales", "Net_Sales"
}


def connect_db():
    if not os.path.isfile(DB_PATH):
        raise FileNotFoundError(
            "Database not found:\n" + DB_PATH
        )

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    columns = {
        row["name"]
        for row in conn.execute(
            "PRAGMA table_info(sales)"
        ).fetchall()
    }

    missing = REQUIRED_COLUMNS - columns

    if missing:
        conn.close()
        raise ValueError(
            "Missing database columns: "
            + ", ".join(sorted(missing))
        )

    return conn


def money(value):
    return "Rs. {:,.2f}".format(value or 0)


def fetch_all(conn, sql, params=()):
    return [
        dict(row)
        for row in conn.execute(sql, params).fetchall()
    ]


def show_table(rows):
    if not rows:
        print("No matching data found.")
        return

    headers = list(rows[0].keys())
    values = [
        [str(row.get(h, "")) for h in headers]
        for row in rows
    ]

    widths = [
        min(
            max(len(headers[i]), *(len(r[i]) for r in values)),
            28
        )
        for i in range(len(headers))
    ]

    print(" | ".join(
        headers[i].ljust(widths[i])
        for i in range(len(headers))
    ))
    print("-+-".join("-" * w for w in widths))

    for row in values:
        print(" | ".join(
            row[i].ljust(widths[i])
            for i in range(len(headers))
        ))


def business_summary(conn):
    row = conn.execute("""
        SELECT
            COALESCE(SUM(Net_Sales), 0) AS sales,
            COUNT(DISTINCT Order_ID) AS orders
        FROM sales
    """).fetchone()

    sales = row["sales"]
    orders = row["orders"]

    print("\nBUSINESS SUMMARY")
    print("-" * 32)
    print("Total net sales:", money(sales))
    print("Distinct orders:", orders)
    print(
        "Average order value:",
        money(sales / orders if orders else 0)
    )

    print("\nTop branch by revenue:")
    show_table(fetch_all(conn, """
        SELECT Branch, ROUND(SUM(Net_Sales), 2) AS Revenue
        FROM sales
        GROUP BY Branch
        ORDER BY Revenue DESC
        LIMIT 1
    """))

    print("\nTop product by revenue:")
    show_table(fetch_all(conn, """
        SELECT Product, ROUND(SUM(Net_Sales), 2) AS Revenue
        FROM sales
        GROUP BY Product
        ORDER BY Revenue DESC
        LIMIT 1
    """))


def monthly_sales(conn):
    rows = fetch_all(conn, """
        SELECT
            SUBSTR(Order_Date, 1, 7) AS Month,
            ROUND(SUM(Net_Sales), 2) AS Revenue
        FROM sales
        GROUP BY SUBSTR(Order_Date, 1, 7)
        ORDER BY Month
    """)

    print("\nMONTHLY REVENUE")
    show_table(rows)

    if len(rows) >= 2:
        previous = rows[-2]["Revenue"]
        current = rows[-1]["Revenue"]

        if previous:
            change = (current - previous) / abs(previous) * 100
            print("Latest month change: {:+.2f}%".format(change))
        else:
            print("Previous month revenue is zero; change unavailable.")


def branch_analysis(conn, lowest=False):
    direction = "ASC" if lowest else "DESC"
    title = "LOWEST" if lowest else "BRANCH REVENUE RANKING"

    print("\n" + title)
    rows = fetch_all(conn, """
        SELECT
            Branch,
            COUNT(DISTINCT Order_ID) AS Orders,
            ROUND(SUM(Net_Sales), 2) AS Revenue
        FROM sales
        GROUP BY Branch
        ORDER BY SUM(Net_Sales) """ + direction)

    show_table(rows)


def product_analysis(conn, units=False, lowest=False):
    direction = "ASC" if lowest else "DESC"

    if units:
        metric = "SUM(Quantity)"
        alias = "Units_Sold"
        title = "PRODUCTS BY UNITS SOLD"
    else:
        metric = "SUM(Net_Sales)"
        alias = "Revenue"
        title = "PRODUCTS BY REVENUE"

    print("\n" + title)

    rows = fetch_all(conn, """
        SELECT
            Product,
            ROUND(""" + metric + """, 2) AS """ + alias + """
        FROM sales
        GROUP BY Product
        ORDER BY """ + metric + " " + direction)

    show_table(rows)


def payment_analysis(conn):
    print("\nPAYMENT METHOD ANALYSIS")
    show_table(fetch_all(conn, """
        SELECT
            Payment_Method,
            COUNT(DISTINCT Order_ID) AS Orders,
            ROUND(SUM(Net_Sales), 2) AS Revenue
        FROM sales
        GROUP BY Payment_Method
        ORDER BY Revenue DESC
    """))


def category_analysis(conn):
    print("\nCATEGORY PERFORMANCE")
    show_table(fetch_all(conn, """
        SELECT
            Category,
            SUM(Quantity) AS Units_Sold,
            ROUND(SUM(Net_Sales), 2) AS Revenue
        FROM sales
        GROUP BY Category
        ORDER BY Revenue DESC
    """))


def discount_analysis(conn):
    print("\nDISCOUNT ANALYSIS")
    show_table(fetch_all(conn, """
        SELECT
            ROUND(SUM(Gross_Sales), 2) AS Gross_Sales,
            ROUND(SUM(Net_Sales), 2) AS Net_Sales,
            ROUND(
                SUM(Gross_Sales) - SUM(Net_Sales), 2
            ) AS Sales_Difference
        FROM sales
    """))
    print(
        "Note: Sales difference is not necessarily "
        "the exact discount cost unless the data "
        "definitions confirm that relationship."
    )


def create_excel_report(conn):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    total = conn.execute("""
        SELECT
            COALESCE(SUM(Net_Sales), 0),
            COUNT(DISTINCT Order_ID)
        FROM sales
    """).fetchone()

    sales = total[0]
    orders = total[1]

    summary.append(["Restaurant Sales Analytics"])
    summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    summary.append(["Total Net Sales", sales])
    summary.append(["Distinct Orders", orders])
    summary.append([
        "Average Order Value",
        sales / orders if orders else 0
    ])

    reports = [
        ("Branch Performance", """
            SELECT Branch,
                   COUNT(DISTINCT Order_ID) AS Orders,
                   SUM(Quantity) AS Units_Sold,
                   SUM(Gross_Sales) AS Gross_Sales,
                   SUM(Net_Sales) AS Revenue
            FROM sales GROUP BY Branch ORDER BY Revenue DESC
        """),
        ("Product Performance", """
            SELECT Product,
                   SUM(Quantity) AS Units_Sold,
                   SUM(Gross_Sales) AS Gross_Sales,
                   SUM(Net_Sales) AS Revenue
            FROM sales GROUP BY Product ORDER BY Revenue DESC
        """),
        ("Monthly Revenue", """
            SELECT SUBSTR(Order_Date, 1, 7) AS Month,
                   SUM(Net_Sales) AS Revenue
            FROM sales
            GROUP BY SUBSTR(Order_Date, 1, 7)
            ORDER BY Month
        """),
        ("Payment Analysis", """
            SELECT Payment_Method,
                   COUNT(DISTINCT Order_ID) AS Orders,
                   SUM(Net_Sales) AS Revenue
            FROM sales GROUP BY Payment_Method ORDER BY Revenue DESC
        """),
        ("Category Analysis", """
            SELECT Category,
                   SUM(Quantity) AS Units_Sold,
                   SUM(Net_Sales) AS Revenue
            FROM sales GROUP BY Category ORDER BY Revenue DESC
        """),
        ("Raw Data", """
            SELECT Order_ID, Order_Date, Branch, Product,
                   Category, Quantity, Unit_Price, Discount,
                   Payment_Method, Gross_Sales, Net_Sales
            FROM sales ORDER BY Order_Date
        """)
    ]

    for sheet_name, sql in reports:
        ws = wb.create_sheet(sheet_name)
        rows = fetch_all(conn, sql)

        if rows:
            ws.append(list(rows[0].keys()))
            for row in rows:
                ws.append(list(row.values()))

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="17365D"
            )
            cell.alignment = Alignment(
                vertical="center"
            )

        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            max_len = max(
                len(str(cell.value or ""))
                for cell in column
            )
            ws.column_dimensions[letter].width = min(
                max_len + 2, 30
            )

    wb.save(REPORT_PATH)

    print("\nEXCEL REPORT CREATED SUCCESSFULLY")
    print("Saved to:", REPORT_PATH)


def recommendations(conn):
    branches = fetch_all(conn, """
        SELECT Branch, SUM(Net_Sales) AS Revenue
        FROM sales GROUP BY Branch ORDER BY Revenue DESC
    """)

    products = fetch_all(conn, """
        SELECT Product,
               SUM(Quantity) AS Units,
               SUM(Net_Sales) AS Revenue
        FROM sales GROUP BY Product ORDER BY Revenue DESC
    """)

    print("\nBUSINESS RECOMMENDATIONS")

    if branches:
        print(
            "- Study the practices at the highest-revenue "
            "branch:", branches[0]["Branch"]
        )
        if len(branches) > 1:
            print(
                "- Review operations at the lowest-revenue "
                "branch:", branches[-1]["Branch"]
            )

    if products:
        print(
            "- Review pricing, availability and margins for "
            "the highest-revenue product:",
            products[0]["Product"]
        )

        lowest = min(products, key=lambda x: x["Revenue"])
        print(
            "- Investigate demand and profitability for:",
            lowest["Product"]
        )

    print(
        "- Revenue alone does not measure profit. "
        "Check food cost, labour and other operating expenses."
    )


def answer_question(conn, question):
    q = re.sub(r"[^a-z0-9\s]", " ", question.lower())
    q = " ".join(q.split())

    if not q:
        print("Please enter a question.")
        return

    if q in ("help", "examples", "what can you do"):
        print("""
EXAMPLE QUESTIONS
- Give me a business summary
- What are total sales?
- Which branch has the highest revenue?
- Which branch performs worst?
- Compare all branches
- Which product sells the most units?
- Show top products by revenue
- Show monthly sales
- Analyse payment methods
- Show category performance
- Analyse discounts
- Give business recommendations
- Generate an Excel report
- Exit
""")
        return

    if q in ("exit", "quit", "stop"):
        raise SystemExit

    if any(x in q for x in (
        "generate an excel report", "create excel report",
        "export excel", "make excel report"
    )):
        create_excel_report(conn)

    elif any(x in q for x in (
        "recommendation", "recommendations",
        "what should we do", "action plan",
        "how can i improve", "how to improve"
    )):
        recommendations(conn)

    elif any(x in q for x in (
        "business summary", "overall summary",
        "overall performance", "give me a summary"
    )):
        business_summary(conn)

    elif any(x in q for x in (
        "monthly", "month wise", "month by month",
        "sales trend", "revenue trend"
    )):
        monthly_sales(conn)

    elif any(x in q for x in (
        "payment", "upi", "cash", "card payment"
    )):
        payment_analysis(conn)

    elif "discount" in q or "gross sales" in q:
        discount_analysis(conn)

    elif "category" in q or "categories" in q:
        category_analysis(conn)

    elif "branch" in q or "branches" in q:
        if any(x in q for x in (
            "lowest", "worst", "least", "weakest"
        )):
            branch_analysis(conn, lowest=True)
        else:
            branch_analysis(conn)

    elif any(x in q for x in (
        "unit", "units sold", "quantity", "most items"
    )):
        product_analysis(conn, units=True)

    elif "product" in q or "products" in q or "item" in q:
        product_analysis(conn)

    elif any(x in q for x in (
        "total sales", "net sales", "total revenue",
        "how much did we sell", "total collection"
    )):
        row = conn.execute("""
            SELECT COALESCE(SUM(Net_Sales), 0) AS Total_Sales
            FROM sales
        """).fetchone()
        print("\nTotal net sales:", money(row["Total_Sales"]))

    elif any(x in q for x in (
        "order count", "number of orders", "total orders"
    )):
        row = conn.execute("""
            SELECT COUNT(DISTINCT Order_ID) AS Orders
            FROM sales
        """).fetchone()
        print("\nTotal distinct orders:", row["Orders"])

    else:
        print("""
I could not confidently identify that question.

Try:
- Give me a business summary
- Which branch has the highest revenue?
- Which product sells the most units?
- Show monthly sales
- Analyse payment methods
- Give business recommendations
- Generate an Excel report

Type 'help' to see examples.
""")


def main():
    print("=" * 44)
    print(" RESTAURANT DATA ANALYST AGENT V2")
    print("=" * 44)

    try:
        conn = connect_db()

        count = conn.execute(
            "SELECT COUNT(*) FROM sales"
        ).fetchone()[0]

        print("Database connected successfully.")
        print("Rows loaded:", count)
        print("Type 'help' for example questions.")
        print("Type 'exit' to close the agent.")

        while True:
            try:
                question = input(
                    "\nAsk your business question: "
                )
                answer_question(conn, question)

            except SystemExit:
                break

            except Exception as error:
                print("Could not complete request:", error)

        conn.close()
        print("Agent closed.")

    except Exception as error:
        print("\nSTARTUP ERROR:", error)
        print("Check the database path and table schema.")


if __name__ == "__main__":
    main()
